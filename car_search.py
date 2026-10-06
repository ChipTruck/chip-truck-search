import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import openpyxl
from openpyxl.cell.cell import MergedCell
import re
import os
from datetime import datetime
from PIL import Image

# ページの設定
st.set_page_config(
    page_title="車両管理＆検索アプリ", 
    page_icon="🚗", 
    layout="wide"
)

# ── スタイル設定 ──
st.markdown('''
<style>
.app-main-title {
    font-size: 1.55rem !important;
    font-weight: bold !important;
    color: #1e88e5;
    margin-top: -10px;
    margin-bottom: 8px;
    white-space: nowrap !important;
}

[data-testid="stDataFrame"] {
    width: 100% !important;
}

a.anchor-link, [data-testid="stHeaderActionElements"] {
    display: none !important;
}

div[data-testid="stTextInput"] input {
    font-size: 1.4rem !important;
    height: 50px !important;
    font-weight: bold !important;
    border: 2px solid #1e88e5 !important;
    border-radius: 8px !important;
    padding: 0 12px !important;
}

div[data-testid="stForm"] {
    border: none !important;
    padding: 0 !important;
}

div[data-testid="stHorizontalBlock"] {
    display: flex !important;
    flex-direction: row !important;
    flex-wrap: nowrap !important;
    align-items: center !important;
    gap: 6px !important;
}
div[data-testid="stHorizontalBlock"] > div {
    flex: 1 1 0% !important;
    min-width: 0 !important;
}
div[data-testid="stHorizontalBlock"] button, div[data-testid="stHorizontalBlock"] a {
    width: 100% !important;
    height: 48px !important;
    font-size: 1.05rem !important;
    font-weight: bold !important;
    border-radius: 8px !important;
    padding: 0 !important;
    display: flex !important;
    align-items: center !important;
    justify-content: center !important;
}

.vehicle-detail-card {
    background-color: #ffffff;
    border: 2px solid #1e88e5;
    border-radius: 14px;
    padding: 16px;
    margin-top: 15px;
    margin-bottom: 20px;
    box-shadow: 0 4px 12px rgba(0,0,0,0.08);
}
.card-title {
    font-size: 1.4rem;
    font-weight: bold;
    color: #1e88e5;
    margin-bottom: 10px;
    border-bottom: 2px solid #e3f2fd;
    padding-bottom: 6px;
}
.card-row {
    font-size: 1.1rem;
    margin: 8px 0;
    color: #222;
}
.card-label {
    font-weight: bold;
    color: #555;
    display: inline-block;
    width: 85px;
}
.fuutai-box {
    background: #eef7ff;
    border: 2px solid #1e88e5;
    border-radius: 10px;
    padding: 10px 14px;
    margin-top: 12px;
    display: flex;
    align-items: center;
    justify-content: space-between;
}
.fuutai-label {
    font-size: 1.25rem;
    font-weight: bold;
    color: #1565c0;
}
.fuutai-val {
    font-size: 2.1rem;
    font-weight: 900;
    color: #0d47a1;
    letter-spacing: 1px;
}
</style>
''', unsafe_allow_html=True)

FILE_PATH = "新_登録車両資料_連動版.xlsx"
SHEET_NAME = "新_登録車両資料"

def safe_str(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ["none", "nan", "<na>"]:
        return ""
    return s

def normalize_text(text):
    s = safe_str(text)
    if not s:
        return ""
    t = s.replace("：", ":").replace(" ", "").replace("　", "")
    return t.strip()

@st.cache_data(ttl=5)
def load_shared_excel_data():
    if not os.path.exists(FILE_PATH):
        return pd.DataFrame(columns=['会社名', '車番', '入力番号', '詳細', '備考', '風体'])

    try:
        raw_df = pd.read_excel(FILE_PATH, sheet_name=SHEET_NAME, header=None)
    except Exception:
        raw_df = pd.read_excel(FILE_PATH, header=None)

    raw_df = raw_df.dropna(how="all").astype(str)
    raw_df = raw_df.replace(r'^\s*$', pd.NA, regex=True)

    processed_rows = []
    current_comp = ""
    
    for _, row in raw_df.iterrows():
        row_vals = [safe_str(val) for val in row.values if safe_str(val)]
        if not row_vals:
            continue

        # 会社見出しの判定：行の最初の値に「:」があるときだけ見出しとみなす
        # （登録日時「2026-10-06 11:46:00」の「:」で車両行が見出し扱いされて消える不具合の対策）
        first_val = row_vals[0]
        if _is_company_header(first_val):
            current_comp = first_val.strip()
            continue
        
        if current_comp and len(row_vals) > 0:
            padded = [safe_str(v) for v in row.values]
            new_row = [current_comp] + padded
            processed_rows.append(new_row)

    max_len = max(len(r) for r in processed_rows) if processed_rows else 2
    base_col_names = ['会社名', '車番', '入力番号', '詳細', '備考', '削除対象', '風体', '登録日時', '削除フラグ']
    
    columns = []
    for i in range(max_len):
        if i < len(base_col_names):
            columns.append(base_col_names[i])
        else:
            columns.append(f"extra_{i}")
    
    padded_rows = []
    for r in processed_rows:
        while len(r) < len(columns):
            r.append("")
        padded_rows.append(r)
        
    df = pd.DataFrame(padded_rows, columns=columns[:len(padded_rows[0])] if padded_rows else columns)
    df = df.dropna(subset=['会社名'])
    
    if '削除対象' in df.columns:
        df = df.drop(columns=['削除対象'])
    if '削除フラグ' in df.columns:
        df = df[df['削除フラグ'] != '1']
        
    for c in ['車番', '入力番号', '詳細', '備考', '風体']:
        if c not in df.columns:
            df[c] = ""
            
    for col in df.columns:
        df[col] = df[col].apply(safe_str)
        
    def get_comp_code(val):
        m = re.match(r'^(\d{1,2})', safe_str(val))
        return int(m.group(1)) if m else 999

    df['_comp_num'] = df['会社名'].apply(get_comp_code)
    df = df.sort_values(by=['_comp_num', '車番'], ascending=[True, True], kind='stable').drop(columns=['_comp_num']).reset_index(drop=True)
    return df

if "search_query" not in st.session_state:
    st.session_state.search_query = ""
if "active_card_key" not in st.session_state:
    st.session_state.active_card_key = None
if "scanned_shaban" not in st.session_state:
    st.session_state.scanned_shaban = ""
if "scanned_detail" not in st.session_state:
    st.session_state.scanned_detail = ""
if "action_notice" not in st.session_state:
    st.session_state.action_notice = ""
if "edit_search_keyword" not in st.session_state:
    st.session_state.edit_search_keyword = ""
if "form_reset_counter" not in st.session_state:
    st.session_state.form_reset_counter = 0

def reset_to_home():
    st.session_state.search_query = ""
    st.session_state.active_card_key = None
    st.session_state.action_notice = ""
    st.session_state.form_reset_counter += 1

# ── Excel操作の補助（結合セル対策） ──
def _is_company_header(text):
    t = safe_str(text)
    if not t:
        return False
    if re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}', t):
        return False
    return (":" in t) or ("：" in t)

def _set_cell(ws, row, col, value):
    """結合セルだった場合は結合を解除してから書き込む"""
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for mr in list(ws.merged_cells.ranges):
            if cell.coordinate in mr:
                ws.unmerge_cells(mr.coord)
                break
    ws.cell(row=row, column=col).value = value

def _insert_rows_safe(ws, idx, amount=1):
    """行を挿入するとき、結合セルの位置も一緒にずらす（openpyxlは自動でずらさないため）"""
    saved = []
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row >= idx:
            saved.append((mr.min_row, mr.min_col, mr.max_row, mr.max_col))
            ws.unmerge_cells(mr.coord)
    ws.insert_rows(idx, amount)
    for r1, c1, r2, c2 in saved:
        if r1 >= idx:
            r1 += amount
            r2 += amount
        else:
            r2 += amount
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

def _delete_row_safe(ws, idx):
    """行を削除するとき、結合セルの位置も一緒にずらす"""
    saved = []
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row >= idx:
            saved.append((mr.min_row, mr.min_col, mr.max_row, mr.max_col))
            ws.unmerge_cells(mr.coord)
    ws.delete_rows(idx)
    for r1, c1, r2, c2 in saved:
        if r1 > idx:
            r1 -= 1
            r2 -= 1
        else:
            r2 -= 1
        if r2 < r1 or (r1 == r2 and c1 == c2):
            continue
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

def _find_vehicle_row(ws, company_str, shaban, input_no):
    """会社見出し行・空行を除いて、対象の車両行を1行だけ探す"""
    norm_c = normalize_text(company_str)
    norm_s = normalize_text(shaban)
    norm_inp = normalize_text(input_no)

    current_comp = ""
    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(row=r, column=1).value
        if _is_company_header(c1):
            current_comp = normalize_text(c1)
            continue
        vals = [normalize_text(ws.cell(row=r, column=c).value) for c in range(2, 6)]
        if not any(vals):
            continue
        c2, c3 = vals[0], vals[1]
        if norm_inp and c3 == norm_inp:
            return r
        if current_comp == norm_c and c2 == norm_s:
            return r
    return None

def insert_vehicle_record(company_str, shaban, input_no, detail, remark, fuutai):
    wb = openpyxl.load_workbook(FILE_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    
    comp_row_idx = None
    target_comp_code = -1
    m_target = re.match(r'^(\d{1,2})[:：]', company_str.strip())
    if m_target:
        target_comp_code = int(m_target.group(1))

    insert_before_row = None
    for r in range(1, ws.max_row + 1):
        v = normalize_text(ws.cell(row=r, column=1).value)
        if not v:
            continue
        if normalize_text(company_str) == v:
            comp_row_idx = r
            break
        m = re.match(r'^(\d{1,2})[:：]', v)
        if m and target_comp_code >= 0:
            c_code = int(m.group(1))
            if c_code > target_comp_code and insert_before_row is None:
                insert_before_row = r

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if comp_row_idx is not None:
        insert_target = comp_row_idx + 1
        while insert_target <= ws.max_row:
            cell_v = ws.cell(row=insert_target, column=1).value
            if _is_company_header(cell_v) and normalize_text(cell_v) != normalize_text(company_str):
                break
            insert_target += 1
        
        _insert_rows_safe(ws, insert_target)
        data_row = insert_target
    else:
        ins_r = insert_before_row if insert_before_row is not None else (ws.max_row + 1)
        _insert_rows_safe(ws, ins_r, amount=2)
        _set_cell(ws, ins_r, 1, str(company_str))
        data_row = ins_r + 1

    _set_cell(ws, data_row, 1, "")
    _set_cell(ws, data_row, 2, str(shaban))
    _set_cell(ws, data_row, 3, str(input_no))
    _set_cell(ws, data_row, 4, str(detail))
    _set_cell(ws, data_row, 5, str(remark))
    _set_cell(ws, data_row, 6, "")
    _set_cell(ws, data_row, 7, str(fuutai))
    _set_cell(ws, data_row, 8, now_str)
        
    wb.save(FILE_PATH)
    st.cache_data.clear()

def delete_vehicle_record(company_str, old_shaban, old_input_no):
    wb = openpyxl.load_workbook(FILE_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    r = _find_vehicle_row(ws, company_str, old_shaban, old_input_no)
    if r is None:
        raise ValueError("削除する車両がExcel内に見つかりませんでした。")
    _delete_row_safe(ws, r)

    wb.save(FILE_PATH)
    st.cache_data.clear()

def update_vehicle_record(old_company, old_shaban, old_input_no, new_comp, new_shaban, new_input_no, new_detail, new_remark, new_fuutai):
    wb = openpyxl.load_workbook(FILE_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    r = _find_vehicle_row(ws, old_company, old_shaban, old_input_no)
    if r is None:
        raise ValueError("更新する車両がExcel内に見つかりませんでした。")

    _set_cell(ws, r, 2, str(new_shaban))
    _set_cell(ws, r, 3, str(new_input_no))
    _set_cell(ws, r, 4, str(new_detail))
    _set_cell(ws, r, 5, str(new_remark))
    _set_cell(ws, r, 7, str(new_fuutai))

    wb.save(FILE_PATH)
    st.cache_data.clear()

def generate_upload_csv(df_source):
    rows = []
    for _, r in df_source.iterrows():
        v_id = safe_str(r.get('入力番号', ''))
        fuutai_raw = safe_str(r.get('風体', ''))
        
        if v_id:
            m_fuutai = re.findall(r'\d+', fuutai_raw)
            fuutai_val = m_fuutai[0] if m_fuutai else "0"
            rows.append({
                'Vehicle ID': v_id,
                'Max weight': 0,
                'Weight': fuutai_val
            })
            
    csv_df = pd.DataFrame(rows, columns=['Vehicle ID', 'Max weight', 'Weight'])
    return csv_df.to_csv(index=False, encoding='utf-8-sig')

df_base = load_shared_excel_data()

# ── ヘッダー ──
st.markdown('<div class="app-main-title">🚗 車両管理＆検索</div>', unsafe_allow_html=True)

col_home, col_dl = st.columns(2)
with col_home:
    st.button("🏠 ホーム", on_click=reset_to_home, use_container_width=True)
with col_dl:
    csv_data = generate_upload_csv(df_base)
    st.download_button(
        label="📥 Upload.csv",
        data=csv_data,
        file_name="Upload.csv",
        mime="text/csv",
        use_container_width=True
    )

tab1, tab2, tab3 = st.tabs(["🔍 検索", "➕ 新規登録", "✏️ 編集・削除"])

# ── 【タブ1】 検索・閲覧 ──
with tab1:
    st.markdown("##### 🔍 車番を入力（例: 8, 14, 1351 など）")

    form_key = f"search_form_{st.session_state.form_reset_counter}"
    with st.form(form_key):
        inp_val = st.text_input(
            "車番を入力", 
            value=st.session_state.search_query,
            placeholder="タップして車番を入力",
            label_visibility="collapsed"
        )
        c_search, c_clear = st.columns(2)
        submit_search = c_search.form_submit_button("🔍 検索実行", type="primary", use_container_width=True)
        submit_clear = c_clear.form_submit_button("✕ クリア", use_container_width=True)

        if submit_search:
            st.session_state.search_query = inp_val.strip()
            st.session_state.active_card_key = None
            st.rerun()

        if submit_clear:
            reset_to_home()
            st.rerun()

    # 検索欄をタップしたらテンキーが出るようにする
    components.html("""
    <script>
    const doc = window.parent.document;
    function setNumericKeyboard() {
        doc.querySelectorAll('input[placeholder="タップして車番を入力"]').forEach(el => {
            if (el.getAttribute('inputmode') !== 'numeric') {
                el.setAttribute('inputmode', 'numeric');
                el.setAttribute('pattern', '[0-9]*');
                el.setAttribute('autocomplete', 'off');
            }
        });
    }
    setNumericKeyboard();
    new MutationObserver(setNumericKeyboard).observe(doc.body, {childList: true, subtree: true});
    </script>
    """, height=0)

    current_search = st.session_state.search_query.strip()
    filtered_df = df_base.copy()
    is_searched = bool(current_search)

    if is_searched:
        if '車番' in filtered_df.columns:
            mask = filtered_df['車番'].astype(str).str.contains(current_search, case=False, na=False)
            filtered_df = filtered_df[mask]

        try:
            if '車番' in filtered_df.columns:
                filtered_df['_sort_num'] = pd.to_numeric(filtered_df['車番'].str.extract(r'(\d+)', expand=False), errors='coerce')
                filtered_df = filtered_df.sort_values(by='_sort_num', ascending=True, na_position='last')
                filtered_df = filtered_df.drop(columns=['_sort_num'])
        except Exception:
            pass

    display_cols = [c for c in ['会社名', '車番', '入力番号', '詳細', '備考', '風体'] if c in filtered_df.columns]
    display_df = filtered_df[display_cols].copy()

    table_df = display_df.copy()
    if '会社名' in table_df.columns:
        table_df['会社名'] = table_df['会社名'].mask(table_df['会社名'] == table_df['会社名'].shift(), '')

    if is_searched:
        st.write(f"検索結果: **{len(filtered_df)}** 行 （検索キー: {current_search}）")
    else:
        st.write(f"全車両一覧: **{len(filtered_df)}** 行")

    st.dataframe(table_df, use_container_width=True, hide_index=True)

    if is_searched and len(filtered_df) > 0:
        st.write("---")
        st.markdown("##### 👆 車両を押すとカードで詳細が出ます")
        
        cols = st.columns(2)
        for idx, (_, row) in enumerate(filtered_df.iterrows()):
            shaban_txt = row.get('車番', '-')
            comp_txt = row.get('会社名', '-')
            btn_label = f"🚗 {shaban_txt} （{comp_txt}）"
            
            if cols[idx % 2].button(btn_label, key=f"car_btn_{idx}", use_container_width=True):
                st.session_state.active_card_key = idx
                st.rerun()

        current_active = st.session_state.get("active_card_key", None)
        if current_active is not None and current_active < len(filtered_df):
            target_row = filtered_df.iloc[current_active]
            c_sh = target_row.get('車番', '-')
            c_cp = target_row.get('会社名', '-')
            c_in = target_row.get('入力番号', '-')
            c_dt = target_row.get('詳細', '-')
            c_rm = target_row.get('備考', '-')
            c_ft = target_row.get('風体', '-')
            
            card_html = (
                '<div class="vehicle-detail-card">'
                f'<div class="card-title">🚗 車番: {c_sh}</div>'
                f'<div class="card-row"><span class="card-label">会社名:</span> <b>{c_cp}</b></div>'
                f'<div class="card-row"><span class="card-label">入力番号:</span> <b style="color: #d32f2f; font-size: 1.55rem; letter-spacing: 1px;">{c_in}</b></div>'
                f'<div class="card-row"><span class="card-label">詳細:</span> {c_dt}</div>'
                f'<div class="card-row"><span class="card-label">備考:</span> <b>{c_rm}</b></div>'
                '<div class="fuutai-box">'
                '<span class="fuutai-label">⚖️ 風体</span>'
                f'<span class="fuutai-val">{c_ft}</span>'
                '</div>'
                '</div>'
            )
            st.markdown(card_html, unsafe_allow_html=True)

# ── 【タブ2】 新規登録 ──
with tab2:
    st.subheader("➕ 新規車両の登録")

    if st.session_state.action_notice:
        st.success(st.session_state.action_notice)
        st.session_state.action_notice = ""
    
    existing_companies = [c for c in df_base['会社名'].dropna().unique().tolist() if safe_str(c)]
    used_numbers = set()
    for c in existing_companies:
        m = re.match(r'^(\d{1,2})[:：]', str(c).strip())
        if m:
            used_numbers.add(int(m.group(1)))
    
    next_avail_num = 0
    while next_avail_num in used_numbers and next_avail_num < 100:
        next_avail_num += 1
    default_2digit = f"{next_avail_num:02d}"

    comp_mode = st.radio("① 会社の指定方法", ["既存の会社から選ぶ", "新しい会社を番号から決める"], horizontal=True)
    
    selected_target_comp = ""
    comp_2digit_prefix = ""

    if comp_mode == "既存の会社から選ぶ" and existing_companies:
        selected_target_comp = st.selectbox("登録先の会社名を選択", existing_companies)
        m = re.match(r'^(\d{1,2})[:：]', str(selected_target_comp).strip())
        if m:
            comp_2digit_prefix = f"{int(m.group(1)):02d}"
    else:
        st.markdown("##### 🏢 新しい会社の番号と名前を決める")
        c_num_col, c_name_col = st.columns([1, 2])
        
        with c_num_col:
            new_comp_code = st.text_input("会社番号(2桁)", value=default_2digit, max_chars=2)
        with c_name_col:
            new_comp_raw_name = st.text_input("会社名（例: 木村木材）")

        if new_comp_code:
            if not new_comp_code.isdigit() or len(new_comp_code) != 2:
                st.warning("会社番号は半角数字2桁（00〜99）で入力してください。")
            elif int(new_comp_code) in used_numbers:
                st.error(f"番号「{new_comp_code}」は既に使われています！別の空き番号を指定してください。")
            else:
                st.info(f"番号「{new_comp_code}」は空いています。利用可能です！")

        if new_comp_code and new_comp_raw_name:
            selected_target_comp = f"{new_comp_code}：{new_comp_raw_name.strip()}"
            comp_2digit_prefix = new_comp_code

    st.write("---")
    with st.expander("カメラ撮影 / 画像アップロードで自動入力する", expanded=False):
        upload_choice = st.radio("入力方法", ["カメラで撮影", "写真をアップロード"], horizontal=True)
        uploaded_image = None
        if upload_choice == "カメラで撮影":
            uploaded_image = st.camera_input("シャッターを押して撮影")
        else:
            uploaded_image = st.file_uploader("写真ファイルを選択", type=["jpg", "jpeg", "png"])

        if uploaded_image is not None:
            st.image(uploaded_image, caption="取り込んだ画像", width=250)
            if st.button("画像から車番・ナンバーを読み取る", use_container_width=True):
                extracted_text = ""
                try:
                    import pytesseract
                    img = Image.open(uploaded_image)
                    extracted_text = pytesseract.image_to_string(img, lang="jpn+eng")
                except Exception:
                    pass
                
                nums = re.findall(r'\b\d{1,4}\b', extracted_text)
                if nums:
                    st.session_state.scanned_shaban = nums[-1]
                    clean_detail = extracted_text.strip().replace(chr(10), " ")
                    st.session_state.scanned_detail = clean_detail
                    st.success(f"読み取り成功！ 車番「{st.session_state.scanned_shaban}」を下に入力しました。")
                else:
                    st.info("画像を受け付けました。下のカードで必要項目を確認・入力してください。")
                st.rerun()

    if selected_target_comp:
        st.markdown(f'''
        <div class="vehicle-detail-card">
            <div class="card-title">📝 【{selected_target_comp}】の新規車両カード</div>
        </div>
        ''', unsafe_allow_html=True)

        with st.form("new_vehicle_form_card"):
            new_shaban = st.text_input(
                "車番 *必須（重複時はA/B等）", 
                value=st.session_state.get("scanned_shaban", "")
            )
            
            calc_input_no = f"{comp_2digit_prefix}{new_shaban}" if comp_2digit_prefix and new_shaban else ""
            new_input_no = st.text_input("入力番号（会社2桁＋車番）", value=calc_input_no)
            
            new_detail = st.text_input(
                "詳細（例: 岐阜302 も 9418）", 
                value=st.session_state.get("scanned_detail", "")
            )
            new_remark = st.text_input("備考（例: 4t車、大型など）")
            new_fuutai = st.text_input("風体")

            submitted = st.form_submit_button("この内容で登録を保存", type="primary")
            if submitted:
                if not new_shaban:
                    st.error("車番を入力してください！")
                else:
                    try:
                        insert_vehicle_record(
                            selected_target_comp,
                            new_shaban,
                            new_input_no,
                            new_detail,
                            new_remark,
                            new_fuutai
                        )
                        st.session_state.scanned_shaban = ""
                        st.session_state.scanned_detail = ""
                        st.session_state.action_notice = f"会社「{selected_target_comp}」に 車番「{new_shaban}」を追加しました！PC・スマホ共に反映されています。"
                        st.toast("登録が完了しました！")
                        st.rerun()
                    except Exception as ex:
                        st.error(f"保存中にエラーが発生しました: {ex}")

# ── 【タブ3】 編集・削除 ──
with tab3:
    st.subheader("✏️ 車両情報の編集・削除")

    if st.session_state.action_notice:
        st.success(st.session_state.action_notice)
        st.session_state.action_notice = ""

    st.write("車番、会社名、入力番号、詳細（ナンバー）のいずれかで検索できます。")

    empty_shaban_count = 0
    for _, r in df_base.iterrows():
        if not safe_str(r.get('車番', '')):
            empty_shaban_count += 1

    if empty_shaban_count > 0:
        if st.button(f"車番が未入力の車両（{empty_shaban_count}台）を抽出する", use_container_width=True):
            st.session_state.edit_search_keyword = "車番未設定"

    edit_search_val = st.text_input(
        "🔍 検索キーワード（例: 8、細川、058、吉田 など）", 
        value=st.session_state.get("edit_search_keyword", ""),
        key="edit_search_input_box"
    )
    st.session_state.edit_search_keyword = edit_search_val

    if edit_search_val:
        s_term = edit_search_val.strip()
        
        matched_indices = []
        if s_term == "車番未設定":
            for idx, r in df_base.iterrows():
                if not safe_str(r.get('車番', '')):
                    matched_indices.append(idx)
        else:
            for idx, r in df_base.iterrows():
                shaban_str = safe_str(r.get('車番', ''))
                comp_str = safe_str(r.get('会社名', ''))
                inp_str = safe_str(r.get('入力番号', ''))
                detail_str = safe_str(r.get('詳細', ''))
                
                if (s_term.lower() in shaban_str.lower() or 
                    s_term.lower() in comp_str.lower() or 
                    s_term.lower() in inp_str.lower() or 
                    s_term.lower() in detail_str.lower()):
                    matched_indices.append(idx)

        matched = df_base.loc[matched_indices] if matched_indices else pd.DataFrame()

        if len(matched) > 0:
            car_choices = []
            for _, r in matched.iterrows():
                c_shaban = safe_str(r.get('車番', '')) or '(車番未設定)'
                c_comp = safe_str(r.get('会社名', ''))
                c_inp = safe_str(r.get('入力番号', ''))
                c_detail = safe_str(r.get('詳細', ''))
                car_choices.append(f"🚗 車番: {c_shaban} | {c_comp} (入力番号: {c_inp} / 詳細: {c_detail})")

            selected_edit_car = st.selectbox("対象の車両を決定してください", car_choices)
            
            edit_idx = car_choices.index(selected_edit_car)
            target_edit_row = matched.iloc[edit_idx]
            target_old_comp = safe_str(target_edit_row.get('会社名', ''))
            target_old_shaban = safe_str(target_edit_row.get('車番', ''))
            target_old_input_no = safe_str(target_edit_row.get('入力番号', ''))

            st.markdown(f'''
            <div class="vehicle-detail-card">
                <div class="card-title">✏️ 車両編集・削除カード: {target_old_shaban or '（車番未設定）'}</div>
            </div>
            ''', unsafe_allow_html=True)

            with st.form("card_edit_form"):
                e_comp = st.text_input("会社名", value=target_old_comp)
                e_shaban = st.text_input("車番（ここに番号を入力）", value=target_old_shaban)
                e_input = st.text_input("入力番号", value=target_old_input_no)
                e_detail = st.text_input("詳細", value=safe_str(target_edit_row.get('詳細', '')))
                e_remark = st.text_input("備考", value=safe_str(target_edit_row.get('備考', '')))
                e_fuutai = st.text_input("風体", value=safe_str(target_edit_row.get('風体', '')))

                c1, c2 = st.columns(2)
                save_clicked = c1.form_submit_button("変更を保存", type="primary")
                del_clicked = c2.form_submit_button("この車両を削除")

                if save_clicked:
                    try:
                        update_vehicle_record(
                            target_old_comp, target_old_shaban, target_old_input_no,
                            e_comp, e_shaban, e_input, e_detail, e_remark, e_fuutai
                        )
                        st.session_state.action_notice = f"車両情報（車番「{e_shaban}」）の内容を更新・保存しました！"
                        st.toast("変更を保存しました！")
                        st.rerun()
                    except Exception as ex:
                        st.error(f"保存中にエラーが発生しました: {ex}")

                if del_clicked:
                    try:
                        delete_vehicle_record(target_old_comp, target_old_shaban, target_old_input_no)
                        st.session_state.action_notice = f"会社「{target_old_comp}」の車両を完全に消去しました！"
                        st.toast("データを消去しました！")
                        st.rerun()
                    except Exception as ex:
                        st.error(f"削除中にエラーが発生しました: {ex}")
        else:
            st.info("該当する車両が見つかりませんでした。別のキーワード（会社名や詳細など）をお試しください。")
